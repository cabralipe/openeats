import io

from openpyxl import Workbook
from django.http import HttpResponse
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from rest_framework import permissions, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.response import Response

from accounts.permissions import CanAccessPnaeModule, scope_queryset_by_municipality

from .models import (
    PnaeAcceptabilityTest,
    PnaeAnnualAction,
    PnaeAnnualBudgetItem,
    PnaeAnnualEvaluationTool,
    PnaeAnnualGoal,
    PnaeAnnualPlan,
    PnaeAnnualPlanItem,
    PnaeAnnualPlanMonthlyExecution,
    PnaeAnnualPlanWorkflowEvent,
    PnaeAnnualScheduleEntry,
)
from .serializers import (
    PnaeAcceptabilityTestSerializer,
    PnaeAnnualActionSerializer,
    PnaeAnnualBudgetItemSerializer,
    PnaeAnnualEvaluationToolSerializer,
    PnaeAnnualGoalSerializer,
    PnaeAnnualPlanItemSerializer,
    PnaeAnnualPlanMonthlyExecutionSerializer,
    PnaeAnnualPlanSerializer,
    PnaeAnnualPlanSummarySerializer,
    PnaeAnnualScheduleEntrySerializer,
    PnaeOperationalSummaryRequestSerializer,
    PnaeWorkflowActionSerializer,
)
from .services import (
    PnaeWorkflowError,
    approve_plan,
    build_dashboard_metrics,
    build_operational_summary,
    ensure_plan_editable,
    generate_delivery_draft_from_plan,
    generate_menu_drafts_from_plan,
    record_workflow_event,
    reject_plan,
    submit_plan_for_review,
    sync_monthly_execution_snapshots,
)

def _scope_plan_queryset(queryset, user):
    queryset = scope_queryset_by_municipality(queryset, user, lookup='school__municipality_id')
    if getattr(user, 'is_pnae_viewer_restricted', False):
        queryset = queryset.exclude(status=PnaeAnnualPlan.Status.DRAFT)
    return queryset


def _scope_plan_related_queryset(queryset, user):
    queryset = scope_queryset_by_municipality(queryset, user, lookup='plan__school__municipality_id')
    if getattr(user, 'is_pnae_viewer_restricted', False):
        queryset = queryset.exclude(plan__status=PnaeAnnualPlan.Status.DRAFT)
    return queryset


class PlanScopedQuerysetMixin:
    def get_queryset(self):
        queryset = super().get_queryset()
        queryset = _scope_plan_related_queryset(queryset, self.request.user)
        plan = self.request.query_params.get('plan')
        if plan:
            queryset = queryset.filter(plan_id=plan)
        return queryset

    def perform_destroy(self, instance):
        try:
            ensure_plan_editable(instance.plan)
        except ValueError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        super().perform_destroy(instance)


class PnaeAnnualPlanViewSet(viewsets.ModelViewSet):
    queryset = PnaeAnnualPlan.objects.select_related(
        'school', 'school__municipality', 'created_by', 'responsible_nutritionist',
        'submitted_by', 'approved_by', 'rejected_by',
    ).prefetch_related(
        'items__recipe__ingredients__supply',
        'goals',
        'actions',
        'schedule_entries',
        'budget_items',
        'evaluation_tools',
        'workflow_events__actor',
        'monthly_executions__last_updated_by',
    )
    permission_classes = [permissions.IsAuthenticated, CanAccessPnaeModule]

    def get_serializer_class(self):
        if self.action == 'list':
            return PnaeAnnualPlanSummarySerializer
        return PnaeAnnualPlanSerializer

    def get_queryset(self):
        queryset = _scope_plan_queryset(super().get_queryset(), self.request.user)
        school = self.request.query_params.get('school')
        year = self.request.query_params.get('year')
        status_value = self.request.query_params.get('status')
        responsible_nutritionist = self.request.query_params.get('responsible_nutritionist')
        municipality = self.request.query_params.get('municipality')
        if school:
            queryset = queryset.filter(school_id=school)
        if year:
            queryset = queryset.filter(year=year)
        if status_value:
            queryset = queryset.filter(status=status_value)
        if responsible_nutritionist:
            queryset = queryset.filter(responsible_nutritionist_id=responsible_nutritionist)
        if municipality:
            queryset = queryset.filter(school__municipality_id=municipality)
        return queryset

    def perform_create(self, serializer):
        extra = {'created_by': self.request.user}
        if 'responsible_nutritionist' not in serializer.validated_data and getattr(self.request.user, 'is_nutritionist', False):
            extra['responsible_nutritionist'] = self.request.user
        plan = serializer.save(**extra)
        record_workflow_event(
            plan,
            action=PnaeAnnualPlanWorkflowEvent.Action.CREATED,
            actor=self.request.user,
            from_status='',
            to_status=plan.status,
            comment='Plano criado.',
        )
        sync_monthly_execution_snapshots(plan)

    def perform_update(self, serializer):
        instance = serializer.instance
        previous_status = instance.status
        try:
            ensure_plan_editable(instance)
        except ValueError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        plan = serializer.save()
        sync_monthly_execution_snapshots(plan)
        event_action = (
            PnaeAnnualPlanWorkflowEvent.Action.ARCHIVED
            if previous_status != plan.status and plan.status == PnaeAnnualPlan.Status.ARCHIVED
            else PnaeAnnualPlanWorkflowEvent.Action.UPDATED
        )
        record_workflow_event(
            plan,
            action=event_action,
            actor=self.request.user,
            from_status=previous_status,
            to_status=plan.status,
            comment='Plano atualizado pelo usuario.',
        )

    def perform_destroy(self, instance):
        try:
            ensure_plan_editable(instance)
        except ValueError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        super().perform_destroy(instance)

    @action(detail=False, methods=['get'], url_path='dashboard')
    def dashboard(self, request):
        return Response(build_dashboard_metrics(self.get_queryset()))

    @action(detail=True, methods=['post'], url_path='submit-review')
    def submit_review(self, request, pk=None):
        plan = self.get_object()
        serializer = PnaeWorkflowActionSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            submit_plan_for_review(plan, request.user, serializer.validated_data.get('comment', ''))
        except PnaeWorkflowError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        return Response(PnaeAnnualPlanSerializer(plan, context={'request': request}).data)

    @action(detail=True, methods=['post'], url_path='approve')
    def approve(self, request, pk=None):
        plan = self.get_object()
        serializer = PnaeWorkflowActionSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            approve_plan(plan, request.user, serializer.validated_data.get('comment', ''))
        except PnaeWorkflowError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        return Response(PnaeAnnualPlanSerializer(plan, context={'request': request}).data)

    @action(detail=True, methods=['post'], url_path='reject')
    def reject(self, request, pk=None):
        plan = self.get_object()
        serializer = PnaeWorkflowActionSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            reject_plan(plan, request.user, serializer.validated_data.get('comment', ''))
        except PnaeWorkflowError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        return Response(PnaeAnnualPlanSerializer(plan, context={'request': request}).data)

    @action(detail=True, methods=['get'], url_path='operational-summary')
    def operational_summary(self, request, pk=None):
        plan = self.get_object()
        serializer = PnaeOperationalSummaryRequestSerializer(data=request.query_params)
        serializer.is_valid(raise_exception=True)
        selected_month = serializer.validated_data.get('month')
        sync_monthly_execution_snapshots(plan)
        return Response(build_operational_summary(plan, selected_month=selected_month))

    @action(detail=True, methods=['post'], url_path='generate-menu-drafts')
    def generate_menu_drafts(self, request, pk=None):
        plan = self.get_object()
        if plan.status != PnaeAnnualPlan.Status.APPROVED:
            raise ValidationError({'detail': 'Os rascunhos operacionais so podem ser gerados para planos aprovados.'})
        serializer = PnaeOperationalSummaryRequestSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        month = serializer.validated_data.get('month')
        if not month:
            raise ValidationError({'month': 'Informe o mes para gerar os cardapios.'})
        result = generate_menu_drafts_from_plan(plan, month, request.user)
        return Response(result, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'], url_path='generate-delivery-draft')
    def generate_delivery_draft(self, request, pk=None):
        plan = self.get_object()
        if plan.status != PnaeAnnualPlan.Status.APPROVED:
            raise ValidationError({'detail': 'A entrega so pode ser gerada para planos aprovados.'})
        serializer = PnaeOperationalSummaryRequestSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        month = serializer.validated_data.get('month')
        if not month:
            raise ValidationError({'month': 'Informe o mes para gerar a entrega.'})
        try:
            result = generate_delivery_draft_from_plan(plan, month, request.user)
        except PnaeWorkflowError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        return Response(result, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['get'], url_path='export-xlsx')
    def export_xlsx(self, request, pk=None):
        plan = self.get_object()
        serializer = PnaeOperationalSummaryRequestSerializer(data=request.query_params)
        serializer.is_valid(raise_exception=True)
        selected_month = serializer.validated_data.get('month')
        operational = build_operational_summary(plan, selected_month=selected_month)
        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = 'Resumo'
        summary_sheet.append(['Plano', plan.title or f'{plan.school.name} {plan.year}'])
        summary_sheet.append(['Escola', plan.school.name])
        summary_sheet.append(['Municipio', plan.school.municipality.name if plan.school.municipality else plan.city])
        summary_sheet.append(['Ano', plan.year])
        summary_sheet.append([])
        summary_sheet.append(['Indicador', 'Valor'])
        for key, value in operational['annual_summary'].items():
            summary_sheet.append([key, value])

        items_sheet = workbook.create_sheet('Operacional')
        items_sheet.append([
            'Mes', 'Etapa', 'Modalidade', 'Refeicao', 'Receita',
            'Ocorrencias', 'Porcoes projetadas', 'Custo estimado',
        ])
        for item in operational['detail']['items']:
            items_sheet.append([
                operational['detail']['month_label'],
                item['education_stage_name'],
                item['education_modality_name'],
                item['meal_type_display'],
                item['recipe_name'],
                item['monthly_occurrences'],
                item['projected_servings'],
                item['estimated_cost'],
            ])

        procurement_sheet = workbook.create_sheet('Compras')
        procurement_sheet.append([
            'Insumo', 'Unidade', 'Quantidade', 'Custo estimado',
            'Saldo escola', 'Saldo central', 'Falta escola', 'Falta central',
        ])
        for item in operational['detail']['procurement']:
            procurement_sheet.append([
                item['supply_name'],
                item['unit'],
                item['qty_needed'],
                item['estimated_cost'],
                item['school_stock_available'],
                item['central_stock_available'],
                item['school_shortage'],
                item['central_shortage'],
            ])

        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="pnae-{plan.id}.xlsx"'
        return response

    @action(detail=True, methods=['get'], url_path='export-pdf')
    def export_pdf(self, request, pk=None):
        plan = self.get_object()
        serializer = PnaeOperationalSummaryRequestSerializer(data=request.query_params)
        serializer.is_valid(raise_exception=True)
        selected_month = serializer.validated_data.get('month')
        operational = build_operational_summary(plan, selected_month=selected_month)
        buffer = io.BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4
        top = height - 40
        pdf.setFont('Helvetica-Bold', 16)
        pdf.drawString(40, top, 'Resumo do Plano PNAE')
        pdf.setFont('Helvetica', 10)
        lines = [
            f'Plano: {plan.title or f"{plan.school.name} {plan.year}"}',
            f'Escola: {plan.school.name}',
            f'Municipio: {plan.school.municipality.name if plan.school.municipality else plan.city}',
            f'Status: {plan.get_status_display()}',
            f'Mes operacional: {operational["selected_month_label"]}',
            f'Porcoes projetadas: {operational["detail"]["summary"]["projected_servings"]}',
            f'Custo estimado: R$ {operational["detail"]["summary"]["estimated_cost"]:.2f}',
            f'Itens com falta na escola: {operational["detail"]["summary"]["supplies_with_shortage"]}',
        ]
        cursor = top - 30
        for line in lines:
            pdf.drawString(40, cursor, line)
            cursor -= 16
        cursor -= 8
        pdf.setFont('Helvetica-Bold', 12)
        pdf.drawString(40, cursor, 'Compras / Estoque')
        cursor -= 20
        pdf.setFont('Helvetica', 9)
        for item in operational['detail']['procurement'][:18]:
            pdf.drawString(
                40,
                cursor,
                f"{item['supply_name']} | {item['qty_needed']} {item['unit']} | falta escola {item['school_shortage']}",
            )
            cursor -= 14
            if cursor < 60:
                pdf.showPage()
                cursor = height - 40
                pdf.setFont('Helvetica', 9)
        pdf.showPage()
        pdf.save()
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="pnae-{plan.id}.pdf"'
        return response


class PnaeAnnualPlanItemViewSet(PlanScopedQuerysetMixin, viewsets.ModelViewSet):
    queryset = PnaeAnnualPlanItem.objects.select_related(
        'plan', 'plan__school', 'plan__school__municipality', 'education_stage',
        'education_modality', 'recipe',
    ).all()
    serializer_class = PnaeAnnualPlanItemSerializer
    permission_classes = [permissions.IsAuthenticated, CanAccessPnaeModule]

    def get_queryset(self):
        queryset = super().get_queryset()
        school = self.request.query_params.get('school')
        month = self.request.query_params.get('month')
        if school:
            queryset = queryset.filter(plan__school_id=school)
        if month:
            queryset = queryset.filter(month=month)
        return queryset

    def perform_create(self, serializer):
        plan = serializer.validated_data['plan']
        try:
            ensure_plan_editable(plan)
        except ValueError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        item = serializer.save()
        sync_monthly_execution_snapshots(item.plan)

    def perform_update(self, serializer):
        plan = serializer.instance.plan
        try:
            ensure_plan_editable(plan)
        except ValueError as exc:
            raise ValidationError({'detail': str(exc)}) from exc
        item = serializer.save()
        sync_monthly_execution_snapshots(item.plan)

    def perform_destroy(self, instance):
        plan = instance.plan
        super().perform_destroy(instance)
        sync_monthly_execution_snapshots(plan)


class PnaeAnnualGoalViewSet(PlanScopedQuerysetMixin, viewsets.ModelViewSet):
    queryset = PnaeAnnualGoal.objects.select_related('plan', 'plan__school').all()
    serializer_class = PnaeAnnualGoalSerializer
    permission_classes = [permissions.IsAuthenticated, CanAccessPnaeModule]


class PnaeAnnualActionViewSet(PlanScopedQuerysetMixin, viewsets.ModelViewSet):
    queryset = PnaeAnnualAction.objects.select_related('plan', 'plan__school').all()
    serializer_class = PnaeAnnualActionSerializer
    permission_classes = [permissions.IsAuthenticated, CanAccessPnaeModule]


class PnaeAnnualScheduleEntryViewSet(PlanScopedQuerysetMixin, viewsets.ModelViewSet):
    queryset = PnaeAnnualScheduleEntry.objects.select_related('plan', 'plan__school').all()
    serializer_class = PnaeAnnualScheduleEntrySerializer
    permission_classes = [permissions.IsAuthenticated, CanAccessPnaeModule]


class PnaeAnnualBudgetItemViewSet(PlanScopedQuerysetMixin, viewsets.ModelViewSet):
    queryset = PnaeAnnualBudgetItem.objects.select_related('plan', 'plan__school').all()
    serializer_class = PnaeAnnualBudgetItemSerializer
    permission_classes = [permissions.IsAuthenticated, CanAccessPnaeModule]


class PnaeAnnualEvaluationToolViewSet(PlanScopedQuerysetMixin, viewsets.ModelViewSet):
    queryset = PnaeAnnualEvaluationTool.objects.select_related('plan', 'plan__school').all()
    serializer_class = PnaeAnnualEvaluationToolSerializer
    permission_classes = [permissions.IsAuthenticated, CanAccessPnaeModule]


class PnaeAnnualPlanMonthlyExecutionViewSet(PlanScopedQuerysetMixin, viewsets.ModelViewSet):
    queryset = PnaeAnnualPlanMonthlyExecution.objects.select_related(
        'plan', 'plan__school', 'last_updated_by',
    ).all()
    serializer_class = PnaeAnnualPlanMonthlyExecutionSerializer
    permission_classes = [permissions.IsAuthenticated, CanAccessPnaeModule]

    def get_queryset(self):
        queryset = super().get_queryset()
        month = self.request.query_params.get('month')
        if month:
            queryset = queryset.filter(month=month)
        return queryset

    def perform_create(self, serializer):
        plan = serializer.validated_data['plan']
        if plan.status != PnaeAnnualPlan.Status.APPROVED:
            raise PermissionDenied('A execucao mensal so pode ser registrada para planos aprovados.')
        sync_monthly_execution_snapshots(plan)
        serializer.save(last_updated_by=self.request.user)

    def perform_update(self, serializer):
        if serializer.instance.plan.status != PnaeAnnualPlan.Status.APPROVED:
            raise PermissionDenied('A execucao mensal so pode ser atualizada para planos aprovados.')
        serializer.save(last_updated_by=self.request.user)


class PnaeAcceptabilityTestViewSet(viewsets.ModelViewSet):
    queryset = PnaeAcceptabilityTest.objects.select_related(
        'school',
        'school__municipality',
        'menu',
        'recipe',
        'previous_test',
        'created_by',
    ).all()
    serializer_class = PnaeAcceptabilityTestSerializer
    permission_classes = [permissions.IsAuthenticated, CanAccessPnaeModule]

    def get_queryset(self):
        queryset = scope_queryset_by_municipality(super().get_queryset(), self.request.user, lookup='school__municipality_id')
        if getattr(self.request.user, 'is_pnae_viewer_restricted', False):
            queryset = queryset.filter(approved=True)
        school = self.request.query_params.get('school')
        method = self.request.query_params.get('method')
        approved = self.request.query_params.get('approved')
        objective = self.request.query_params.get('objective')
        date_from = self.request.query_params.get('date_from')
        date_to = self.request.query_params.get('date_to')
        if school:
            queryset = queryset.filter(school_id=school)
        if method:
            queryset = queryset.filter(method=method)
        if objective:
            queryset = queryset.filter(objective=objective)
        if approved in {'true', 'false'}:
            queryset = queryset.filter(approved=approved == 'true')
        if date_from:
            queryset = queryset.filter(test_date__gte=date_from)
        if date_to:
            queryset = queryset.filter(test_date__lte=date_to)
        return queryset

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)

    @action(detail=False, methods=['get'], url_path='dashboard')
    def dashboard(self, request):
        queryset = self.get_queryset()
        method_counts = {
            method: queryset.filter(method=method).count()
            for method, _ in PnaeAcceptabilityTest.Method.choices
        }
        approved_count = queryset.filter(approved=True).count()
        failed_count = queryset.filter(approved=False).count()
        pending_retest = queryset.filter(approved=False, next_retest_date__isnull=False).count()
        low_adhesion = queryset.filter(
            adhesion_classification__in=[
                PnaeAcceptabilityTest.AdhesionClassification.LOW,
                PnaeAcceptabilityTest.AdhesionClassification.VERY_LOW,
            ]
        ).count()
        return Response({
            'total_tests': queryset.count(),
            'approved_tests': approved_count,
            'failed_tests': failed_count,
            'pending_retest': pending_retest,
            'low_adhesion_tests': low_adhesion,
            'tests_by_method': method_counts,
            'latest_tests': PnaeAcceptabilityTestSerializer(queryset[:5], many=True, context={'request': request}).data,
        })
